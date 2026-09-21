import os
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document as DocxDocument
from pydantic import AnyHttpUrl, SecretStr
from pymilvus import MilvusClient  # type: ignore[import-untyped]
from sqlalchemy import select

from agent_api.core.config import ModelProvider, Settings
from agent_api.knowledge.application.ingestion import (
    IngestionPipeline,
    IngestionPipelineError,
    ParentChildChunker,
)
from agent_api.knowledge.application.management import SqlAlchemyKnowledgeManagementService
from agent_api.knowledge.infrastructure.database.models import Document, DocumentVersion
from agent_api.knowledge.infrastructure.database.runtime import DatabaseRuntime
from agent_api.knowledge.infrastructure.docling.parser import DoclingParser
from agent_api.knowledge.infrastructure.jobs.repository import SqlAlchemyIngestionRepository
from agent_api.knowledge.infrastructure.milvus.index import MilvusChunkIndex
from agent_api.knowledge.infrastructure.storage.local import LocalFileStorage
from tests.test_docling_adapter import minimal_text_pdf
from tests.test_knowledge_management_integration import cleanup_knowledge_base

pytestmark = [
    pytest.mark.knowledge_integration,
    pytest.mark.skipif(
        os.getenv("RUN_KNOWLEDGE_INTEGRATION") != "1",
        reason="set RUN_KNOWLEDGE_INTEGRATION=1 to use local knowledge containers",
    ),
]

DATABASE_URL = "postgresql+asyncpg://agent:agent@127.0.0.1:5432/agent"


class DeterministicEmbedding:
    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[float((index % 7) + 1) / 10 for index in range(1536)] for _ in texts]


def docx_bytes() -> bytes:
    output = BytesIO()
    document = DocxDocument()
    document.add_heading("Integration guide", level=1)
    document.add_paragraph("Install the integration safely.")
    document.save(output)
    return output.getvalue()


def settings_for_e2e(storage_root: Path) -> Settings:
    return Settings(
        model_provider=ModelProvider.OPENAI,
        openai_base_url=AnyHttpUrl("https://example.test/v1"),
        openai_api_key=SecretStr("chat-secret"),
        openai_model="chat-model",
        knowledge_enabled=True,
        database_url=SecretStr(DATABASE_URL),
        storage_root=storage_root,
        milvus_uri="http://127.0.0.1:19530",
        milvus_collection="knowledge_chunks",
        embedding_base_url=AnyHttpUrl("https://embedding.example.test/v1"),
        embedding_api_key=SecretStr("embedding-secret"),
        embedding_model="embedding-model",
        embedding_dimension=1536,
        _env_file=None,  # type: ignore[call-arg]
    )


async def test_four_formats_reach_ready_and_corrupt_pdf_fails(tmp_path: Path) -> None:
    settings = settings_for_e2e(tmp_path / "uploads")
    database = DatabaseRuntime(DATABASE_URL)
    storage = LocalFileStorage(settings.storage_root)
    management = SqlAlchemyKnowledgeManagementService(database.session_factory, storage, settings)
    milvus = MilvusClient(uri=settings.milvus_uri)
    index = MilvusChunkIndex(milvus, settings.milvus_collection)
    repository = SqlAlchemyIngestionRepository(database.session_factory, storage)
    pipeline = IngestionPipeline(
        repository=repository,
        parser=DoclingParser(),
        chunker=ParentChildChunker(max_child_characters=500),
        embedder=DeterministicEmbedding(),
        index=index,
    )
    knowledge_base_id: UUID | None = None
    indexed_versions: list[UUID] = []
    try:
        knowledge_base = await management.create_knowledge_base(name="四格式验收", description="")
        knowledge_base_id = knowledge_base["id"]
        fixtures = [
            ("guide.pdf", "application/pdf", minimal_text_pdf("Knowledge integration")),
            (
                "guide.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                docx_bytes(),
            ),
            ("guide.md", "text/markdown", "# 中文标题\n正文内容".encode()),
            ("guide.txt", "text/plain", "纯文本知识".encode()),
        ]
        accepted_documents: list[UUID] = []
        for filename, mime_type, content in fixtures:
            accepted = await management.upload_document(
                knowledge_base_id, filename, mime_type, content
            )
            await pipeline.run(accepted["job_id"])
            accepted_documents.append(accepted["document_id"])

        corrupt = await management.upload_document(
            knowledge_base_id,
            "corrupt.pdf",
            "application/pdf",
            b"%PDF-1.7\ncorrupt",
        )
        with pytest.raises(IngestionPipelineError, match="document_parse_failed"):
            await pipeline.run(corrupt["job_id"])

        async with database.session_factory() as session:
            ready_documents = list(
                await session.scalars(select(Document).where(Document.id.in_(accepted_documents)))
            )
            failed_document = await session.get(Document, corrupt["document_id"])
            indexed_versions = [
                document.active_version_id
                for document in ready_documents
                if document.active_version_id is not None
            ]
            failed_version = await session.scalar(
                select(DocumentVersion).where(DocumentVersion.document_id == corrupt["document_id"])
            )

        assert len(ready_documents) == 4
        assert all(document.status == "ready" for document in ready_documents)
        indexed_counts = [await index.count_version(version_id) for version_id in indexed_versions]
        assert all(count > 0 for count in indexed_counts)
        assert failed_document is not None and failed_document.status == "failed"
        assert failed_version is not None and failed_version.error_code == "document_parse_failed"
    finally:
        for version_id in indexed_versions:
            await index.delete_version(version_id)
        milvus.close()
        if knowledge_base_id is not None:
            await cleanup_knowledge_base(database, knowledge_base_id)
        await database.close()
