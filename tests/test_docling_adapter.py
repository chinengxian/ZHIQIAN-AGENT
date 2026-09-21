from pathlib import Path

import pytest
from docx import Document

from agent_api.knowledge.infrastructure.docling.parser import DoclingParser, DocumentParseError

FIXTURES = Path(__file__).parent / "fixtures" / "documents"


def test_docling_parser_preserves_markdown_heading_path() -> None:
    sections = DoclingParser().parse(FIXTURES / "guide.md")

    assert [section.heading_path for section in sections] == [
        ("产品指南",),
        ("产品指南", "安装"),
    ]
    assert "欢迎使用" in sections[0].text
    assert "执行安装命令" in sections[1].text


def test_docling_parser_supports_utf8_text_with_location_fallback() -> None:
    sections = DoclingParser().parse(FIXTURES / "guide.txt")

    assert len(sections) == 1
    assert sections[0].heading_path == ()
    assert sections[0].page_start is None
    assert "纯文本知识" in sections[0].text


def test_docling_parser_maps_corrupt_input_to_stable_error() -> None:
    with pytest.raises(DocumentParseError, match="document_parse_failed"):
        DoclingParser().parse(FIXTURES / "corrupt.pdf")


def minimal_text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


def test_docling_parser_supports_pdf_page_and_docx_heading(tmp_path: Path) -> None:
    pdf_path = tmp_path / "guide.pdf"
    pdf_path.write_bytes(minimal_text_pdf("Knowledge guide"))
    docx_path = tmp_path / "guide.docx"
    document = Document()
    document.add_heading("Product guide", level=1)
    document.add_paragraph("Install the product safely.")
    document.save(docx_path)

    pdf_sections = DoclingParser().parse(pdf_path)
    docx_sections = DoclingParser().parse(docx_path)

    assert any("Knowledge guide" in section.text for section in pdf_sections)
    assert any(section.page_start == 1 for section in pdf_sections)
    assert any(section.heading_path == ("Product guide",) for section in docx_sections)
